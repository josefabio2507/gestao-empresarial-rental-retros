"""Importacao explicita de parcelas historicas, sem regerar parcelamentos."""
from calendar import monthrange
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import BadZipFile, ZipFile
import unicodedata

from sqlalchemy import update
from sqlalchemy.exc import IntegrityError

from app.extensions import db
from app.models import CentroCusto, FinanceiroCartaoCredito, FinanceiroCartaoFatura, FinanceiroContaPagarTitulo, FinanceiroImportacaoCartao
from app.services.financeiro_contas_pagar_service import calcular_datas_fatura, recalcular_fatura

COLUNAS = [
    "ID legado da parcela", "ID legado da compra", "Cartão (nome cadastrado)",
    "Final do cartão", "Fornecedor", "Descrição", "Documento", "Data da compra",
    "Parcela nº", "Total de parcelas", "Valor da parcela (R$)",
    "Competência da fatura", "Vencimento da fatura", "Centro de custo",
    "Situação conferida", "Observações",
]
MAX_LINHAS = 10000
TAMANHO_LOTE = 100
PAUSA_SEGUNDOS = 10


def texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def normalizado(valor):
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", texto(valor)) if not unicodedata.combining(c)).lower().split())


def _data(valor, campo, competencia=False, epoch=None):
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        from openpyxl.utils.datetime import from_excel, WINDOWS_EPOCH
        try:
            valor = from_excel(valor, epoch=epoch or WINDOWS_EPOCH)
        except (ValueError, OverflowError) as exc:
            raise ValueError(f"{campo}: data inválida.") from exc
    if isinstance(valor, datetime):
        valor = valor.date()
    if not isinstance(valor, date):
        formatos = ["%m/%Y", "%Y-%m"] if competencia else ["%d/%m/%Y", "%Y-%m-%d"]
        for formato in formatos:
            try:
                valor = datetime.strptime(texto(valor), formato).date()
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"{campo}: data inválida ou ausente.")
    return (valor.replace(day=1) if competencia else valor).isoformat()


def _numero(valor):
    if isinstance(valor, bool):
        raise ValueError("Número inválido.")
    bruto = texto(valor).replace("R$", "").replace(" ", "")
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    try:
        numero = Decimal(bruto)
        if not numero.is_finite():
            raise ValueError("Número inválido.")
        return numero
    except InvalidOperation as exc:
        raise ValueError("Número inválido ou ausente.") from exc


def ler_planilha(arquivo):
    from openpyxl import load_workbook

    if Path(arquivo.filename or "").suffix.lower() != ".xlsx":
        raise ValueError("Selecione uma planilha .xlsx.")
    conteudo = arquivo.read(10 * 1024 * 1024 + 1)
    if len(conteudo) > 10 * 1024 * 1024:
        raise ValueError("A planilha deve ter até 10 MB.")
    try:
        with ZipFile(BytesIO(conteudo)) as zipfile:
            if sum(item.file_size for item in zipfile.infolist()) > 60 * 1024 * 1024:
                raise ValueError("Planilha excede o tamanho permitido após descompressão.")
        workbook = load_workbook(BytesIO(conteudo), read_only=True, data_only=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Não foi possível ler a planilha Excel.") from exc
    try:
        if "Parcelas" not in workbook.sheetnames:
            raise ValueError("A planilha precisa conter a aba Parcelas.")
        sheet = workbook["Parcelas"]
        if sheet.max_row and sheet.max_row > MAX_LINHAS + 1:
            raise ValueError(f"Limite de {MAX_LINHAS} linhas excedido. Remova linhas extras da aba Parcelas.")
        cabecalho = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        indices = {normalizado(v): i for i, v in enumerate(cabecalho) if v is not None}
        obrigatorias = [0, 3, 4, 5, 7, 8, 9, 10, 12]
        faltantes = [COLUNAS[i] for i in obrigatorias if normalizado(COLUNAS[i]) not in indices]
        if faltantes:
            raise ValueError("Colunas ausentes: " + ", ".join(faltantes))
        linhas, chaves = [], set()
        for numero, cells in enumerate(sheet.iter_rows(min_row=2, max_row=MAX_LINHAS + 2), 2):
            valores = [cells[indices[normalizado(c)]].value if normalizado(c) in indices and indices[normalizado(c)] < len(cells) else None for c in COLUNAS]
            if not any(v is not None and texto(v) for v in valores):
                continue
            if numero > MAX_LINHAS + 1:
                raise ValueError(f"Limite de {MAX_LINHAS} linhas excedido.")
            linha = {"linha": numero, "id_legado": texto(valores[0]), "final": texto(valores[3]).zfill(4),
                     "fornecedor": texto(valores[4]), "descricao": texto(valores[5]), "documento": texto(valores[6]),
                     "id_compra": texto(valores[1]), "nome_cartao": texto(valores[2]),
                     "centro": texto(valores[13]), "observacoes": texto(valores[15]), "erros_base": []}
            try:
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("Substitua as fórmulas da linha pelos valores calculados.")
                for campo, tamanho in [("id_legado", 80), ("fornecedor", 180), ("descricao", 220)]:
                    if not linha[campo] or len(linha[campo]) > tamanho:
                        raise ValueError(f"{campo}: obrigatório, com até {tamanho} caracteres.")
                if len(linha["documento"]) > 80:
                    raise ValueError("Documento: limite de 80 caracteres.")
                if not linha["final"].isascii() or not linha["final"].isdigit() or len(linha["final"]) != 4 or not texto(valores[3]):
                    raise ValueError("Final do cartão: informe os quatro dígitos.")
                linha["chave"] = f"cartao:{linha['final']}:{linha['id_legado']}"
                if linha["chave"] in chaves:
                    raise ValueError("ID legado repetido para o mesmo cartão nesta planilha.")
                chaves.add(linha["chave"])
                linha["compra"] = _data(valores[7], "Data da compra", epoch=workbook.epoch)
                linha["vencimento"] = _data(valores[12], "Vencimento", epoch=workbook.epoch)
                linha["competencia"] = _data(valores[11] or date.fromisoformat(linha["vencimento"]), "Competência", True, epoch=workbook.epoch)
                parcela, total = _numero(valores[8]), _numero(valores[9])
                if parcela != parcela.to_integral_value() or total != total.to_integral_value() or not 1 <= parcela <= total <= 9999:
                    raise ValueError("Número da parcela ou total de parcelas inválido.")
                linha["parcela"], linha["total_parcelas"] = int(parcela), int(total)
                valor = _numero(valores[10]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                linha["valor"] = str(valor)
                if valor <= 0:
                    raise ValueError("Desconto/crédito ou valor zero: separado para conferência; não será lançado como compra.")
                if valor >= Decimal("10000000000"):
                    raise ValueError("Valor excede o limite de uma parcela.")
                situacao = normalizado(valores[14])
                if situacao not in ("", "pago", "em aberto", "aberto", "a vencer", "vencido", "aguardando conferencia"):
                    raise ValueError("Situação desconhecida. Use Pago, Em aberto ou Aguardando conferência.")
                linha["pago"] = situacao == "pago"
                linha["conferencia"] = situacao == "aguardando conferencia"
            except (ValueError, InvalidOperation) as exc:
                linha["erros_base"].append(str(exc))
            linhas.append(linha)
        if not linhas:
            raise ValueError("Nenhuma parcela encontrada na aba Parcelas.")
        return linhas
    finally:
        workbook.close()


def cartoes_disponiveis():
    # Cartoes inativos podem ter historico de compras legitimo.
    return FinanceiroCartaoCredito.query.order_by(FinanceiroCartaoCredito.nome).all()


def mapear_automaticamente(linhas):
    cartoes = cartoes_disponiveis()
    resultado = {}
    for final in {r["final"] for r in linhas}:
        candidatos = [c for c in cartoes if c.ultimos_4_digitos == final]
        if len(candidatos) == 1:
            resultado[final] = candidatos[0].id
    return resultado


def validar(linhas, mapeamento):
    cartoes = {c.id: c for c in cartoes_disponiveis()}
    centros = CentroCusto.query.all()
    resultado = []
    chaves = [r.get("chave") for r in linhas if r.get("chave")]
    existentes = {}
    for inicio in range(0, len(chaves), 400):
        existentes.update(dict(db.session.query(FinanceiroContaPagarTitulo.chave_legado_cartao, FinanceiroContaPagarTitulo.id).filter(FinanceiroContaPagarTitulo.chave_legado_cartao.in_(chaves[inicio:inicio + 400])).all()))
    for original in linhas:
        linha = dict(original, erros=list(original["erros_base"]), avisos=[])
        cartao = cartoes.get(mapeamento.get(linha["final"]))
        if not cartao:
            linha["erros"].append("Selecione o cartão cadastrado correspondente a este final.")
        else:
            linha["cartao_id"] = cartao.id
            linha["cartao"] = cartao.identificacao_segura
            if linha.get("vencimento") and date.fromisoformat(linha["vencimento"]).day != min(cartao.dia_vencimento, monthrange(date.fromisoformat(linha["vencimento"]).year, date.fromisoformat(linha["vencimento"]).month)[1]):
                linha["avisos"].append("Vencimento da parcela preservado; a fatura mensal usa o vencimento do cartão.")
        linha["centro_id"] = None
        if linha["centro"]:
            candidatos = [c for c in centros if normalizado(c.nome) == normalizado(linha["centro"])]
            if len(candidatos) != 1:
                linha["erros"].append("Centro de custo não encontrado ou ambíguo.")
            else:
                linha["centro_id"] = candidatos[0].id
        linha["titulo_existente"] = existentes.get(linha.get("chave"))
        resultado.append(linha)
    return resultado


def criar_importacao(arquivo, usuario):
    linhas = ler_planilha(arquivo)
    mapeamento = mapear_automaticamente(linhas)
    job = FinanceiroImportacaoCartao(id=str(uuid4()), usuario_id=usuario.id,
        arquivo_nome=Path(arquivo.filename).name[:255], dados=validar(linhas, mapeamento), mapeamento=mapeamento)
    db.session.add(job)
    db.session.commit()
    return job


def resumo(job):
    linhas = job.dados
    aptas = [r for r in linhas if not r["erros"] and not r.get("titulo_existente")]
    return {"total": len(linhas), "aptas": len(aptas), "pagas": sum(r.get("pago", False) for r in aptas),
        "abertas": sum(not r.get("pago", False) and not r.get("conferencia", False) for r in aptas),
        "duplicadas": sum(bool(r.get("titulo_existente")) for r in linhas),
        "pendencias": sum(bool(r["erros"]) and not r.get("titulo_existente") for r in linhas),
        "valor": sum((Decimal(r["valor"]) for r in aptas), Decimal("0.00")),
        "importados": job.importados, "ignorados": job.ignorados, "pendentes": job.pendentes,
        "processados": job.cursor, "concluido": job.concluido}


def atualizar_mapeamento(job, formulario):
    if job.iniciado:
        raise ValueError("A importação já começou; o vínculo dos cartões não pode mais ser alterado.")
    mapeamento = {}
    for final in {r["final"] for r in job.dados}:
        valor = formulario.get("cartao_" + final, "")
        if valor:
            try:
                mapeamento[final] = int(valor)
            except ValueError as exc:
                raise ValueError("Cartão inválido.") from exc
    dados = validar(job.dados, mapeamento)
    atualizado = db.session.execute(update(FinanceiroImportacaoCartao).where(
        FinanceiroImportacaoCartao.id == job.id, FinanceiroImportacaoCartao.iniciado.is_(False),
    ).values(mapeamento=mapeamento, dados=dados).execution_options(synchronize_session=False)).rowcount
    if not atualizado:
        db.session.rollback()
        raise ValueError("A importação já começou. Atualize a página para acompanhar.")
    db.session.commit()


def _importar_linha(linha, usuario_id, cache=None):
    cache = cache if cache is not None else {"cartoes": {}, "faturas": {}}
    if FinanceiroContaPagarTitulo.query.filter_by(chave_legado_cartao=linha["chave"]).first():
        return None
    # Serializa criacao/recalculo de faturas entre lotes de importacoes distintas.
    cartao = cache["cartoes"].get(linha["cartao_id"])
    if not cartao:
        cartao = FinanceiroCartaoCredito.query.filter_by(id=linha["cartao_id"]).with_for_update().one()
        cache["cartoes"][cartao.id] = cartao
    competencia = date.fromisoformat(linha["competencia"])
    chave_fatura = (cartao.id, competencia)
    fatura = cache["faturas"].get(chave_fatura)
    if not fatura:
        fatura = FinanceiroCartaoFatura.query.filter_by(cartao_credito_id=cartao.id, competencia=competencia).with_for_update().first()
    if fatura and fatura.status == "Cancelada":
        raise ValueError("Fatura cancelada: reabra ou confira a fatura antes de importar esta parcela.")
    if not fatura:
        fechamento, vencimento = calcular_datas_fatura(cartao, competencia.year, competencia.month)
        fatura = FinanceiroCartaoFatura(cartao_credito_id=cartao.id, competencia=competencia,
            data_fechamento=fechamento, data_vencimento=vencimento, status="Aberta",
            criado_por_usuario_id=usuario_id, atualizado_por_usuario_id=usuario_id)
        db.session.add(fatura)
        db.session.flush()
    cache["faturas"][chave_fatura] = fatura
    valor = Decimal(linha["valor"])
    vencimento = date.fromisoformat(linha["vencimento"])
    pago = valor if linha["pago"] else Decimal("0.00")
    observacoes = [linha["observacoes"]] if linha["observacoes"] else []
    if linha["id_compra"]:
        observacoes.append("ID da compra no legado: " + linha["id_compra"])
    if linha["pago"]:
        observacoes.append("Pagamento informado no legado; data e comprovante não informados.")
    titulo = FinanceiroContaPagarTitulo(
        id_legado=linha["id_legado"], chave_legado_cartao=linha["chave"],
        fornecedor_nome_snapshot=linha["fornecedor"], descricao=linha["descricao"], numero_documento=linha["documento"] or None,
        origem_lancamento="Legado", tipo_pagamento="Cartao de Credito", forma_pagamento="Cartao de Credito",
        cartao_credito_id=cartao.id, fatura_cartao_id=fatura.id, competencia_fatura_cartao=competencia,
        competencia=competencia, data_emissao=date.fromisoformat(linha["compra"]), data_compra_cartao=date.fromisoformat(linha["compra"]),
        data_vencimento=vencimento, valor_original=valor, valor_pago=pago, valor_pago_legado=pago,
        parcela_numero=linha["parcela"], total_parcelas=linha["total_parcelas"], centro_custo_id=linha["centro_id"],
        status="Pago" if linha["pago"] else "Aguardando conferencia" if linha["conferencia"] else "Vencido" if vencimento < date.today() else "A vencer",
        observacoes="\n".join(observacoes) or None, criado_por_usuario_id=usuario_id, atualizado_por_usuario_id=usuario_id)
    db.session.add(titulo)
    db.session.flush()
    return fatura.id


def processar_lote(job_id, usuario_id, cursor):
    # UPDATE condicional funciona tambem no SQLite: um pedido repetido nunca avanca dois lotes.
    agora = datetime.now(timezone.utc).replace(tzinfo=None)
    adquirido = db.session.execute(update(FinanceiroImportacaoCartao).where(
        FinanceiroImportacaoCartao.id == job_id, FinanceiroImportacaoCartao.usuario_id == usuario_id,
        FinanceiroImportacaoCartao.cursor == cursor, FinanceiroImportacaoCartao.concluido.is_(False),
        (FinanceiroImportacaoCartao.proximo_lote_em.is_(None)) | (FinanceiroImportacaoCartao.proximo_lote_em <= agora),
    ).values(iniciado=True).execution_options(synchronize_session=False)).rowcount
    job = db.session.get(FinanceiroImportacaoCartao, job_id, populate_existing=True)
    if not job or job.usuario_id != usuario_id:
        db.session.rollback()
        raise ValueError("Importação não encontrada.")
    if not adquirido:
        db.session.rollback()
        return resumo(job)
    linhas = [dict(r) for r in job.dados]
    fim = min(cursor + TAMANHO_LOTE, len(linhas))
    cache = {"cartoes": {}, "faturas": {}}
    faturas_alteradas = set()
    for linha in linhas[cursor:fim]:
        if linha.get("titulo_existente"):
            job.ignorados += 1
            continue
        if linha["erros"]:
            job.pendentes += 1
            continue
        try:
            with db.session.begin_nested():
                fatura_id = _importar_linha(linha, usuario_id, cache)
            if fatura_id:
                job.importados += 1
                linha["resultado"] = "Importado"
                faturas_alteradas.add(fatura_id)
            else:
                job.ignorados += 1
                linha["resultado"] = "Já importado"
        except ValueError as exc:
            linha["erros"] = [str(exc)]
            job.pendentes += 1
        except IntegrityError:
            # Outra importacao pode ter confirmado o mesmo ID durante este lote.
            if FinanceiroContaPagarTitulo.query.filter_by(chave_legado_cartao=linha["chave"]).first():
                job.ignorados += 1
                linha["resultado"] = "Já importado"
            else:
                raise
    for fatura_id in faturas_alteradas:
        recalcular_fatura(db.session.get(FinanceiroCartaoFatura, fatura_id))
    job.dados = linhas
    job.cursor = fim
    job.concluido = fim == len(linhas)
    job.proximo_lote_em = datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(seconds=PAUSA_SEGUNDOS)
    db.session.commit()
    return resumo(job)
