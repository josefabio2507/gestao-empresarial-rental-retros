import os
from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.services.financeiro_contas_pagar_service import calcular_saldo_titulo
from app.utils.datas import agora_brasil


COR_AZUL = "082C52"
COR_AZUL_CLARO = "EAF2F8"
COR_AMARELO = "F4B400"
COR_CINZA = "F5F7FA"
COR_TEXTO = "17324D"
COR_BORDA = "D9E2EC"
COR_TEXTO_SECUNDARIO = "68788A"


def _data_formatada(valor):
    if not valor:
        return "-"
    if hasattr(valor, "strftime"):
        return valor.strftime("%d/%m/%Y")
    try:
        return datetime.strptime(str(valor), "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return str(valor)


def _periodo(filtros):
    inicio = filtros.get("vencimento_inicio") or filtros.get("data_inicio")
    fim = filtros.get("vencimento_fim") or filtros.get("data_fim")
    if inicio or fim:
        return f"{_data_formatada(inicio) if inicio else '-'} a {_data_formatada(fim) if fim else '-'}"
    return "Todos os vencimentos"


def _aplicar_borda(intervalo, cor=COR_BORDA):
    borda = Border(
        left=Side(style="thin", color=cor),
        right=Side(style="thin", color=cor),
        top=Side(style="thin", color=cor),
        bottom=Side(style="thin", color=cor),
    )
    for linha in intervalo:
        for celula in linha:
            celula.border = borda


def gerar_excel_titulos(titulos, filtros):
    """Gera o relatório de títulos em Excel com identidade visual do Financeiro."""
    titulos_ordenados = sorted(
        titulos,
        key=lambda titulo: (titulo.valor_original or 0, titulo.id or 0),
        reverse=True,
    )
    gerado_em = agora_brasil()

    workbook = Workbook()
    planilha = workbook.active
    planilha.title = "Títulos a Pagar"
    planilha.sheet_view.showGridLines = False
    planilha.freeze_panes = "A11"
    planilha.sheet_properties.pageSetUpPr.fitToPage = True
    planilha.page_setup.orientation = "landscape"
    planilha.page_setup.fitToWidth = 1
    planilha.page_setup.fitToHeight = 0
    planilha.print_title_rows = "10:10"
    planilha.sheet_properties.tabColor = COR_AZUL

    larguras = {"A": 16, "B": 31, "C": 10, "D": 36, "E": 19, "F": 24, "G": 18}
    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura

    planilha.merge_cells("A1:E1")
    planilha.merge_cells("A2:E2")
    planilha.merge_cells("A3:E3")
    planilha.merge_cells("F1:G4")
    for linha in planilha[1:4]:
        for celula in linha:
            celula.fill = PatternFill("solid", fgColor=COR_AZUL)
    for linha in range(1, 5):
        for coluna in range(1, 8):
            planilha.cell(linha, coluna).fill = PatternFill("solid", fgColor=COR_AZUL)
    planilha.row_dimensions[1].height = 18
    planilha.row_dimensions[2].height = 28
    planilha.row_dimensions[3].height = 19
    planilha.row_dimensions[4].height = 12

    planilha["A1"] = "FINANCEIRO | CONTAS A PAGAR"
    planilha["A1"].font = Font(name="Arial", size=9, color="D8E6F3")
    planilha["A2"] = "Relatório de Títulos a Pagar"
    planilha["A2"].font = Font(name="Arial", size=18, bold=True, color="FFFFFF")
    planilha["A3"] = "Pagamentos programados por período"
    planilha["A3"].font = Font(name="Arial", size=9, color="D8E6F3")
    for referencia in ("A1", "A2", "A3"):
        planilha[referencia].alignment = Alignment(vertical="center")

    caminho_logo = os.path.join(
        os.path.dirname(os.path.dirname(__file__)), "static", "img", "logo-rental-retros.png"
    )
    if os.path.exists(caminho_logo):
        logo = Image(caminho_logo)
        logo.width = 116
        logo.height = 78
        planilha.add_image(logo, "F1")

    planilha.merge_cells("A6:C6")
    planilha.merge_cells("A7:C7")
    planilha.merge_cells("D6:E6")
    planilha.merge_cells("D7:E7")
    planilha.merge_cells("F6:G6")
    planilha.merge_cells("F7:G7")
    planilha["A6"] = "Período selecionado"
    planilha["A7"] = _periodo(filtros)
    planilha["D6"] = "Gerado em"
    planilha["D7"] = gerado_em
    planilha["D7"].number_format = "dd/mm/yyyy hh:mm"
    planilha["F6"] = "Títulos"
    planilha["F7"] = len(titulos_ordenados)
    for linha in range(6, 8):
        for coluna in range(1, 8):
            celula = planilha.cell(linha, coluna)
            celula.fill = PatternFill("solid", fgColor=COR_CINZA)
            celula.font = Font(name="Arial", size=9, bold=linha == 6, color=COR_TEXTO)
            celula.alignment = Alignment(vertical="center", horizontal="right" if coluna >= 6 else "left")
    _aplicar_borda(planilha["A6:G7"])
    planilha.row_dimensions[6].height = 20
    planilha.row_dimensions[7].height = 22

    planilha.merge_cells("A9:G9")
    planilha["A9"] = "Títulos previstos para pagamento"
    planilha["A9"].font = Font(name="Arial", size=11, bold=True, color=COR_AZUL)

    colunas = ["Data prevista", "Fornecedor", "ID", "Descrição", "Documento", "Forma de pagamento", "Valor"]
    for indice, titulo_coluna in enumerate(colunas, start=1):
        celula = planilha.cell(10, indice, titulo_coluna)
        celula.fill = PatternFill("solid", fgColor=COR_AZUL)
        celula.font = Font(name="Arial", size=9, bold=True, color="FFFFFF")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    planilha.row_dimensions[10].height = 29

    primeira_linha_dados = 11
    for indice_linha, titulo_item in enumerate(titulos_ordenados, start=primeira_linha_dados):
        valores = [
            titulo_item.data_vencimento,
            titulo_item.fornecedor_nome_snapshot or "-",
            titulo_item.id,
            titulo_item.descricao or "-",
            titulo_item.numero_documento or "-",
            titulo_item.forma_pagamento or "-",
            float(titulo_item.valor_original or 0),
        ]
        for indice_coluna, valor in enumerate(valores, start=1):
            celula = planilha.cell(indice_linha, indice_coluna, valor)
            celula.font = Font(name="Arial", size=9, color=COR_TEXTO)
            celula.alignment = Alignment(
                horizontal="right" if indice_coluna in (3, 7) else "left",
                vertical="top",
                wrap_text=indice_coluna in (2, 4, 6),
            )
        planilha.cell(indice_linha, 1).number_format = "dd/mm/yyyy"
        planilha.cell(indice_linha, 7).number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        planilha.row_dimensions[indice_linha].height = 30

    ultima_linha_tabela = primeira_linha_dados + len(titulos_ordenados) - 1
    if not titulos_ordenados:
        planilha.cell(primeira_linha_dados, 1, "Nenhum título encontrado para os filtros informados.")
        planilha.merge_cells(start_row=primeira_linha_dados, start_column=1, end_row=primeira_linha_dados, end_column=7)
        planilha.cell(primeira_linha_dados, 1).alignment = Alignment(horizontal="center", vertical="center")

    tabela = Table(displayName="TitulosAPagar", ref=f"A10:G{max(10, ultima_linha_tabela)}")
    tabela.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    planilha.add_table(tabela)

    linha_totais = max(primeira_linha_dados, ultima_linha_tabela) + 3
    total_original = sum((titulo.valor_original or 0 for titulo in titulos_ordenados), 0)
    total_pago = sum((titulo.valor_pago or 0 for titulo in titulos_ordenados), 0)
    saldo_aberto = sum((calcular_saldo_titulo(titulo) or 0 for titulo in titulos_ordenados), 0)
    for deslocamento, (rotulo, valor) in enumerate(
        (("Subtotal dos títulos", total_original), ("Total pago", total_pago), ("Saldo em aberto", saldo_aberto))
    ):
        linha = linha_totais + deslocamento
        planilha.merge_cells(start_row=linha, start_column=5, end_row=linha, end_column=6)
        planilha.cell(linha, 5, rotulo)
        planilha.cell(linha, 7, float(valor))
        planilha.cell(linha, 5).font = Font(name="Arial", size=9, bold=True, color=COR_TEXTO)
        planilha.cell(linha, 7).font = Font(name="Arial", size=9, color=COR_TEXTO)
        planilha.cell(linha, 7).number_format = 'R$ #,##0.00;[Red]-R$ #,##0.00'
        planilha.cell(linha, 7).alignment = Alignment(horizontal="right")
    for coluna in range(5, 8):
        planilha.cell(linha_totais, coluna).border = Border(top=Side(style="medium", color=COR_AMARELO))
        planilha.cell(linha_totais + 2, coluna).border = Border(bottom=Side(style="thin", color=COR_BORDA))

    linha_nota = linha_totais + 4
    planilha.merge_cells(start_row=linha_nota, start_column=1, end_row=linha_nota, end_column=7)
    planilha.cell(linha_nota, 1, f"Relatório gerado em {gerado_em.strftime('%d/%m/%Y %H:%M')}. Valores apresentados em reais.")
    planilha.cell(linha_nota, 1).font = Font(name="Arial", size=8, italic=True, color=COR_TEXTO_SECUNDARIO)
    planilha.oddFooter.center.text = "Rental Retros | Financeiro | Uso interno"
    planilha.oddFooter.right.text = "Página &P de &N"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


def nome_arquivo_titulos_excel():
    return f"relatorio_titulos_a_pagar_{agora_brasil().strftime('%Y%m%d_%H%M%S')}.xlsx"
